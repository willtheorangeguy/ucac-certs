from datetime import datetime, timedelta
from io import BytesIO

import pytest
from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from lss_report.awards import add_years
from lss_report.grid import build_grid
from lss_report.models import MemberRecord, ReportData
from lss_report.web.app import create_app
from lss_report.web.auth import SESSION_COOKIE, Auth
from lss_report.web.repository import ScanRepository, StaffRepository, rows_from_scan
from lss_report.web.scans import TIMEZONE, Verification

PREFERRED = "Robin Preferred"
SOCIETY = "Robert Rivers"


@pytest.fixture
def named_staff(settings, database, monkeypatch):
    monkeypatch.setenv("DISABLE_SCHEDULER", "1")
    generated = datetime.now(TIMEZONE)
    expiry = generated.date() + timedelta(days=14)
    repo = StaffRepository(database)
    member = repo.add(
        name=PREFERRED, society_name=SOCIETY, member_code="RRV001", supervisor=True,
        email="robin@example.org", locker_number="007", boot_size="9.5",
        fit_test_date=add_years(expiry, -2).isoformat(),
        cartridge_date=add_years(expiry, -5).isoformat(),
    )
    scans = ScanRepository(database)
    grid = build_grid(ReportData(generated, [MemberRecord(PREFERRED, "RRV001", source_name=SOCIETY)]))
    scan_id = scans.start(triggered_by="test")
    scans.store(scan_id, grid, {"RRV001": member.id})
    with database.write() as connection:
        connection.execute(
            "INSERT INTO notification_log (staff_id, column_code, expiry_date, threshold, channel, sent_at)"
            " VALUES (?, 'Fit test', ?, 14, 'email', ?)",
            (member.id, expiry.isoformat(), generated.isoformat()),
        )
    app = create_app(settings, database)
    with TestClient(app, follow_redirects=False) as client:
        client.cookies.set(SESSION_COOKIE, Auth(database, settings).create_session("manager@example.org"))
        yield client, repo, member, scan_id


def test_preferred_name_is_used_in_pages_dialogs_and_reminder_history(named_staff):
    client, _, member, _ = named_staff
    for path in ("/", "/lockers", "/reminders"):
        text = client.get(path).text
        assert PREFERRED in text
        assert SOCIETY not in text
    soup = BeautifulSoup(client.get("/staff").text, "html.parser")
    assert soup.select_one(f"#edit-{member.id} h2").get_text() == f"Edit {PREFERRED}"
    assert PREFERRED in soup.select_one(f"#files-{member.id} h2").get_text()
    assert SOCIETY in soup.get_text()  # Still available for comparison and adoption.
    reminders = BeautifulSoup(client.get("/reminders").text, "html.parser")
    assert all(PREFERRED in tr.get_text() for tr in reminders.select("tbody tr"))


@pytest.mark.parametrize("path,builder", [
    ("/export.xlsx", "build_workbook"),
    ("/workplace-first-aiders.xlsx", "build_first_aiders_workbook"),
    ("/lockers.xlsx", "build_lockers_workbook"),
    ("/export.pdf", "build_pdf"),
    ("/workplace-first-aiders.pdf", "build_first_aiders_pdf"),
    ("/lockers.pdf", "build_lockers_pdf"),
])
def test_every_export_uses_the_preferred_name(named_staff, monkeypatch, path, builder):
    from lss_report.web import app

    client, _, _, _ = named_staff
    original_builder = getattr(app, builder)
    rendered_names = []

    def capture(report, output):
        rendered_names.extend(row.name for row in report.rows)
        original_builder(report, output)

    monkeypatch.setattr(app, builder, capture)
    response = client.get(path)
    assert response.status_code == 200
    assert rendered_names == [PREFERRED]
    if path.endswith(".xlsx"):
        sheet = load_workbook(BytesIO(response.content)).active
        names = [cell.value for row in sheet.iter_rows(min_col=1, max_col=1) for cell in row]
        assert PREFERRED in names
        assert SOCIETY not in names
    else:
        assert response.content.startswith(b"%PDF-")


def test_adopting_society_name_changes_the_preferred_name_everywhere(named_staff):
    client, repo, member, _ = named_staff
    response = client.post(f"/staff/{member.id}/adopt-name")
    assert response.headers["location"] == "/staff"
    assert repo.get(member.id).name == SOCIETY
    for path in ("/", "/lockers", "/reminders"):
        text = client.get(path).text
        assert SOCIETY in text
        assert PREFERRED not in text
    workbook = load_workbook(BytesIO(client.get("/export.xlsx").content))
    assert SOCIETY in [row[0].value for row in workbook.active.iter_rows()]


def test_editing_preferred_name_updates_reports_without_changing_society_name(named_staff):
    client, repo, member, _ = named_staff
    client.post(f"/staff/{member.id}/edit", data={"name": "New Preferred"})
    assert repo.get(member.id).name == "New Preferred"
    assert repo.get(member.id).society_name == SOCIETY
    text = client.get("/").text
    assert "New Preferred" in text and SOCIETY not in text


def test_name_changes_keep_scan_errors_attached_to_the_correct_staff(database):
    repo = StaffRepository(database)
    member = repo.add(name=PREFERRED, society_name=SOCIETY, member_code="RRV001")
    record = MemberRecord(PREFERRED, "RRV001", source_name=SOCIETY, error="Member ID was not found.")
    generated = datetime.now(TIMEZONE)
    scans = ScanRepository(database)
    scan_id = scans.start(triggered_by="test")
    scans.store(scan_id, build_grid(ReportData(generated, [record])), {"RRV001": member.id})
    repo.update(member.id, actor="test", name="New Preferred")
    row = rows_from_scan(database, scan_id)[0]
    assert row["staff"].display_name == "New Preferred"
    assert row["error"] == "Member ID was not found."


@pytest.mark.parametrize("action", ["add", "edit"])
def test_red_cross_entry_uses_society_name_for_matching(named_staff, monkeypatch, action):
    client, repo, member, _ = named_staff
    calls = []
    monkeypatch.setattr("lss_report.web.app.verify_member_code",
                        lambda *args: Verification(ok=True, society_name=SOCIETY))

    def verify(number, name, **kwargs):
        calls.append((number, name))
        return Verification(ok=True)

    monkeypatch.setattr("lss_report.web.app.verify_red_cross_number", verify)
    fields = {"name": PREFERRED, "red_cross_fa_number": "103575156"}
    if action == "add":
        fields["member_code"] = "NEW002"
        path = "/staff"
    else:
        path = f"/staff/{member.id}/edit"
    assert client.post(path, data=fields).headers["location"] == "/staff"
    assert calls == [("103575156", SOCIETY)]
    assert repo.get(member.id).name == PREFERRED
