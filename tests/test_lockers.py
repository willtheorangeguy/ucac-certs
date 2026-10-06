from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.platypus import Table

from lss_report import theme
from lss_report.excel import build_lockers_workbook
from lss_report.lockers import LOCKER_HEADINGS, LockerExpiry, LockerReport, LockerRow
from lss_report.pdf import build_lockers_pdf

AS_OF = datetime(2026, 10, 6, 7)


@pytest.fixture
def report():
    def expiry(value):
        return LockerExpiry.for_date(value, AS_OF.date())

    return LockerReport(AS_OF, (
        LockerRow("Robin Rivers", "007", expiry(date(2026, 10, 5)),
                  expiry(date(2026, 11, 5)), "9.5"),
        LockerRow("Sam & Sue <Staff>", "=1+1", expiry(None),
                  expiry(date(2026, 11, 6)), "10", away=True),
    ))


def test_locker_workbook_preserves_details_dates_colours_and_away_section(report, tmp_path):
    output = tmp_path / "lockers.xlsx"
    build_lockers_workbook(report, output)
    sheet = load_workbook(output)["Lockers"]
    assert tuple(cell.value for cell in sheet[6]) == LOCKER_HEADINGS
    assert sheet["A7"].value == "Robin Rivers"
    assert sheet["B7"].value == "007"
    assert sheet["E7"].value == "9.5"
    assert sheet["C7"].value == datetime(2026, 10, 5)
    assert sheet["C7"].number_format == "yyyy-mm-dd"
    assert sheet["C7"].fill.fgColor.rgb[-6:] == theme.EXPIRED
    assert sheet["D7"].fill.fgColor.rgb[-6:] == theme.EXPIRING
    assert sheet["A8"].value == "Away"
    assert sheet["A9"].value == "Sam & Sue <Staff>"
    assert sheet["B9"].value == "=1+1" and sheet["B9"].data_type == "s"
    assert sheet["C9"].value is None
    assert sheet["C9"].fill.fgColor.rgb[-6:] == theme.MISSING
    assert sheet["C9"].font.color.rgb[-6:] == "FFFFFF"
    assert sheet["D9"].fill.patternType is None


def test_locker_pdf_contains_five_columns_and_matching_colours(report, tmp_path, monkeypatch):
    from lss_report import pdf

    tables = []
    real_build = pdf.SimpleDocTemplate.build

    def capture_build(document, story, *args, **kwargs):
        tables.extend(item for item in story if isinstance(item, Table))
        return real_build(document, story, *args, **kwargs)

    monkeypatch.setattr(pdf.SimpleDocTemplate, "build", capture_build)
    output = tmp_path / "lockers.pdf"
    build_lockers_pdf(report, output)
    assert output.read_bytes().startswith(b"%PDF-")
    table = tables[0]
    assert tuple(table._cellvalues[0]) == LOCKER_HEADINGS
    assert table._cellvalues[1][0][0].getPlainText() == "Robin Rivers"
    assert table._cellvalues[1][1][0].getPlainText() == "007"
    assert table._cellvalues[1][2:4] == ["2026-10-05", "2026-11-05"]
    assert table._cellvalues[2][0] == "Away"
    backgrounds = {start: fill for _, start, _, fill in table._bkgrndcmds}
    assert backgrounds[(2, 1)] == colors.HexColor(f"#{theme.EXPIRED}")
    assert backgrounds[(3, 1)] == colors.HexColor(f"#{theme.EXPIRING}")
    assert backgrounds[(2, 3)] == colors.HexColor(f"#{theme.MISSING}")
    assert (3, 3) not in backgrounds


@pytest.mark.parametrize("builder,suffix", [(build_lockers_workbook, "xlsx"), (build_lockers_pdf, "pdf")])
def test_empty_locker_report_downloads_are_valid(builder, suffix, tmp_path):
    output = tmp_path / f"empty.{suffix}"
    builder(LockerReport(AS_OF, ()), output)
    if suffix == "xlsx":
        assert tuple(cell.value for cell in load_workbook(output)["Lockers"][6]) == LOCKER_HEADINGS
    else:
        assert output.read_bytes().startswith(b"%PDF-")


def test_large_locker_pdf_paginates(report, tmp_path):
    output = tmp_path / "large.pdf"
    build_lockers_pdf(LockerReport(AS_OF, report.rows * 50), output)
    raw = output.read_bytes()
    assert raw.count(b"/Type /Page") - raw.count(b"/Type /Pages") > 1
